"""Validate and load the Excel business workbooks into analytics RAW tables."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import sys
from zipfile import BadZipFile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

ROOT = Path(__file__).resolve().parents[2]
BUSINESS_DIR = ROOT / "data" / "business"
LOGGER = logging.getLogger("excel")


@dataclass(frozen=True)
class WorkbookSpec:
    file_name: str
    sheet_name: str
    headers: tuple[str, ...]
    table: str
    key_column: str
    numeric_columns: tuple[str, ...]
    required_text: tuple[str, ...] = ()
    date_columns: tuple[str, ...] = ()


SPECS = (
    WorkbookSpec(
        "marketing_campaigns.xlsx",
        "marketing_campaigns",
        ("campaign_id", "game_id", "campaign", "budget", "start_date", "end_date"),
        "raw_marketing_campaigns",
        "campaign_id",
        ("budget",),
        ("campaign",),
        ("start_date", "end_date"),
    ),
    WorkbookSpec(
        "game_targets.xlsx",
        "game_targets",
        (
            "game_id",
            "target_revenue",
            "target_players",
            "target_sessions",
            "target_rating",
        ),
        "raw_game_targets",
        "game_id",
        ("target_revenue", "target_players", "target_sessions", "target_rating"),
    ),
    WorkbookSpec(
        "budget.xlsx",
        "budget",
        ("campaign_id", "budget", "actual_spend"),
        "raw_campaign_budget",
        "campaign_id",
        ("budget", "actual_spend"),
    ),
)


class WorkbookValidationError(ValueError):
    """Raised when an Excel workbook does not match the documented contract."""


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=BUSINESS_DIR,
        help="Directory containing the three required .xlsx files.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Reload all three workbook snapshots even if their file hashes match.",
    )
    return parser.parse_args(argv)


def _integer(value: Any, field: str, file_name: str, row_number: int) -> int:
    if isinstance(value, bool):
        raise WorkbookValidationError(
            f"{file_name} row {row_number}: {field} must be an integer."
        )
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise WorkbookValidationError(
        f"{file_name} row {row_number}: {field} must be an integer."
    )


def _decimal(value: Any, field: str, file_name: str, row_number: int) -> Decimal:
    if value is None or isinstance(value, bool):
        raise WorkbookValidationError(
            f"{file_name} row {row_number}: {field} must be a number."
        )
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise WorkbookValidationError(
            f"{file_name} row {row_number}: {field} must be a number."
        ) from None
    if not number.is_finite():
        raise WorkbookValidationError(
            f"{file_name} row {row_number}: {field} must be finite."
        )
    return number


def _business_date(value: Any, field: str, file_name: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise WorkbookValidationError(
        f"{file_name} row {row_number}: {field} must be an Excel date or ISO date."
    )


def _normalize_row(
    spec: WorkbookSpec, raw_values: Sequence[Any], row_number: int
) -> dict[str, Any] | None:
    if all(value is None or (isinstance(value, str) and not value.strip()) for value in raw_values):
        return None
    if len(raw_values) != len(spec.headers):
        raise WorkbookValidationError(
            f"{spec.file_name} row {row_number}: expected {len(spec.headers)} cells, "
            f"found {len(raw_values)}."
        )
    row = dict(zip(spec.headers, raw_values))
    for field in spec.required_text:
        value = row[field]
        if not isinstance(value, str) or not value.strip():
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: {field} is required."
            )
        row[field] = value.strip()

    id_columns = {"campaign_id", "game_id"}
    for field in id_columns.intersection(row):
        row[field] = _integer(row[field], field, spec.file_name, row_number)
        if row[field] <= 0:
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: {field} must be positive."
            )
    for field in spec.numeric_columns:
        row[field] = (
            _integer(row[field], field, spec.file_name, row_number)
            if field in {"target_players", "target_sessions"}
            else _decimal(row[field], field, spec.file_name, row_number)
        )
    for field in spec.date_columns:
        row[field] = _business_date(row[field], field, spec.file_name, row_number)

    if spec.file_name == "marketing_campaigns.xlsx":
        if row["budget"] < 0:
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: budget cannot be negative."
            )
        if row["end_date"] < row["start_date"]:
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: end_date precedes start_date."
            )
    elif spec.file_name == "game_targets.xlsx":
        if row["target_revenue"] < 0:
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: target_revenue cannot be negative."
            )
        if row["target_players"] < 0 or row["target_sessions"] < 0:
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: player and session targets "
                "cannot be negative."
            )
        if not Decimal("0") <= row["target_rating"] <= Decimal("5"):
            raise WorkbookValidationError(
                f"{spec.file_name} row {row_number}: target_rating must be 0-5."
            )
    elif row["budget"] < 0 or row["actual_spend"] < 0:
        raise WorkbookValidationError(
            f"{spec.file_name} row {row_number}: budget and actual_spend "
            "cannot be negative."
        )
    return row


def read_workbook(path: Path, spec: WorkbookSpec) -> tuple[list[dict[str, Any]], int]:
    if not path.is_file():
        raise WorkbookValidationError(f"Required workbook is missing: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise WorkbookValidationError(
            f"Cannot read {path.name} as an .xlsx workbook ({type(error).__name__})."
        ) from None
    try:
        if spec.sheet_name not in workbook.sheetnames:
            raise WorkbookValidationError(
                f"{path.name} is missing worksheet '{spec.sheet_name}'."
            )
        sheet = workbook[spec.sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            raise WorkbookValidationError(f"{path.name} is empty.") from None
        headers = tuple(
            value.strip() if isinstance(value, str) else value for value in raw_headers
        )
        if headers != spec.headers:
            raise WorkbookValidationError(
                f"{path.name} headers must be exactly: {', '.join(spec.headers)}."
            )
        normalized: list[dict[str, Any]] = []
        blank_rows = 0
        for row_number, values in enumerate(rows, start=2):
            row = _normalize_row(spec, values, row_number)
            if row is None:
                blank_rows += 1
            else:
                normalized.append(row)
        return normalized, blank_rows
    finally:
        workbook.close()


def deduplicate_rows(
    rows: Iterable[dict[str, Any]], spec: WorkbookSpec
) -> tuple[list[dict[str, Any]], int]:
    unique: dict[int, dict[str, Any]] = {}
    duplicates = 0
    for row in rows:
        key = row[spec.key_column]
        existing = unique.get(key)
        if existing is None:
            unique[key] = row
        elif existing == row:
            duplicates += 1
        else:
            raise WorkbookValidationError(
                f"{spec.file_name} has conflicting rows for "
                f"{spec.key_column}={key}."
            )
    return list(unique.values()), duplicates


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_hashes(connection: psycopg.Connection[Any]) -> dict[str, str]:
    with connection.cursor() as cursor:
        cursor.execute("SELECT file_name, sha256 FROM raw.business_file_loads")
        return dict(cursor.fetchall())


def _validate_relations(
    connection: psycopg.Connection[Any],
    datasets: dict[str, list[dict[str, Any]]],
) -> None:
    campaigns = {row["campaign_id"]: row for row in datasets["marketing_campaigns.xlsx"]}
    budgets = datasets["budget.xlsx"]
    for budget_row in budgets:
        campaign = campaigns.get(budget_row["campaign_id"])
        if campaign is None:
            raise WorkbookValidationError(
                "budget.xlsx contains campaign_id "
                f"{budget_row['campaign_id']} that is absent from marketing_campaigns.xlsx."
            )
        if campaign["budget"] != budget_row["budget"]:
            raise WorkbookValidationError(
                f"Budget mismatch for campaign_id {budget_row['campaign_id']} "
                "between the campaign and budget workbooks."
            )
    missing_budgets = set(campaigns) - {row["campaign_id"] for row in budgets}
    if missing_budgets:
        raise WorkbookValidationError(
            "budget.xlsx is missing campaign IDs: "
            + ", ".join(str(value) for value in sorted(missing_budgets)[:10])
        )

    game_ids = {row["game_id"] for rows in datasets.values() for row in rows if "game_id" in row}
    game_ids.update(row["game_id"] for row in datasets["game_targets.xlsx"])
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT game_id FROM raw.games WHERE game_id = ANY(%s)",
            (list(game_ids),),
        )
        available_ids = {row[0] for row in cursor.fetchall()}
    missing_games = game_ids - available_ids
    if missing_games:
        raise WorkbookValidationError(
            "Excel files reference game IDs absent from analytics raw.games: "
            + ", ".join(str(value) for value in sorted(missing_games)[:10])
        )


def _replace_snapshot(
    connection: psycopg.Connection[Any],
    datasets: dict[str, list[dict[str, Any]]],
    hashes: dict[str, str],
    row_counts: dict[str, int],
) -> None:
    loaded_at = datetime.now(timezone.utc)
    with connection.cursor() as cursor:
        cursor.execute("DELETE FROM raw.raw_campaign_budget")
        cursor.execute("DELETE FROM raw.raw_marketing_campaigns")
        cursor.execute("DELETE FROM raw.raw_game_targets")

        campaigns = datasets["marketing_campaigns.xlsx"]
        if campaigns:
            cursor.executemany(
                """
                INSERT INTO raw.raw_marketing_campaigns
                    (campaign_id, game_id, campaign, budget, start_date, end_date,
                     source_file, loaded_at)
                VALUES (%(campaign_id)s, %(game_id)s, %(campaign)s, %(budget)s,
                        %(start_date)s, %(end_date)s, 'marketing_campaigns.xlsx',
                        %(loaded_at)s)
                """,
                [
                    {**row, "loaded_at": loaded_at}
                    for row in campaigns
                ],
            )
        targets = datasets["game_targets.xlsx"]
        if targets:
            cursor.executemany(
                """
                INSERT INTO raw.raw_game_targets
                    (game_id, target_revenue, target_players, target_sessions,
                     target_rating, source_file, loaded_at)
                VALUES (%(game_id)s, %(target_revenue)s, %(target_players)s,
                        %(target_sessions)s, %(target_rating)s,
                        'game_targets.xlsx', %(loaded_at)s)
                """,
                [{**row, "loaded_at": loaded_at} for row in targets],
            )
        budgets = datasets["budget.xlsx"]
        if budgets:
            cursor.executemany(
                """
                INSERT INTO raw.raw_campaign_budget
                    (campaign_id, budget, actual_spend, source_file, loaded_at)
                VALUES (%(campaign_id)s, %(budget)s, %(actual_spend)s,
                        'budget.xlsx', %(loaded_at)s)
                """,
                [{**row, "loaded_at": loaded_at} for row in budgets],
            )
        cursor.executemany(
            """
            INSERT INTO raw.business_file_loads (file_name, sha256, row_count, loaded_at)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (file_name) DO UPDATE
            SET sha256 = EXCLUDED.sha256,
                row_count = EXCLUDED.row_count,
                loaded_at = EXCLUDED.loaded_at
            """,
            [
                (name, hashes[name], row_counts[name], loaded_at)
                for name in hashes
            ],
        )


def load_files(
    connection: psycopg.Connection[Any], input_dir: Path, force: bool = False
) -> dict[str, Any]:
    datasets: dict[str, list[dict[str, Any]]] = {}
    hashes: dict[str, str] = {}
    report: dict[str, Any] = {"files": {}}
    for spec in SPECS:
        path = input_dir / spec.file_name
        rows, blank_rows = read_workbook(path, spec)
        rows, duplicate_rows = deduplicate_rows(rows, spec)
        if not rows:
            raise WorkbookValidationError(
                f"{spec.file_name} has no business records; refusing to replace "
                "the RAW snapshot with an empty workbook."
            )
        datasets[spec.file_name] = rows
        hashes[spec.file_name] = _sha256(path)
        report["files"][spec.file_name] = {
            "rows": len(rows),
            "blank_rows_skipped": blank_rows,
            "identical_duplicates_skipped": duplicate_rows,
            "sha256": hashes[spec.file_name],
        }

    _validate_relations(connection, datasets)
    previous_hashes = _load_hashes(connection)
    unchanged = all(previous_hashes.get(name) == sha for name, sha in hashes.items())
    if unchanged and not force:
        report["status"] = "skipped_unchanged"
        return report

    row_counts = {name: len(rows) for name, rows in datasets.items()}
    _replace_snapshot(connection, datasets, hashes, row_counts)
    report["status"] = "loaded_snapshot"
    return report


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    load_dotenv(ROOT / ".env")
    args = parse_args(argv)
    if not os.environ.get("POSTGRES_USER") or not os.environ.get("POSTGRES_PASSWORD"):
        LOGGER.error("POSTGRES_USER and POSTGRES_PASSWORD are required in .env.")
        return 2
    try:
        with psycopg.connect(
            host=os.environ.get("ANALYTICS_HOST", "localhost"),
            port=int(os.environ.get("ANALYTICS_PORT", "5433")),
            dbname=os.environ.get("ANALYTICS_DB", "gaming_analytics"),
            user=os.environ["POSTGRES_USER"],
            password=os.environ["POSTGRES_PASSWORD"],
            connect_timeout=10,
        ) as connection:
            schema_path = ROOT / "sql" / "004_excel_business_schema.sql"
            with connection.cursor() as cursor:
                cursor.execute(schema_path.read_text(encoding="utf-8"))
            connection.commit()
            report = load_files(connection, args.input_dir, args.force)
    except (OSError, psycopg.Error, WorkbookValidationError, RuntimeError) as error:
        LOGGER.error("%s", error)
        return 1
    print(json.dumps(report, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
