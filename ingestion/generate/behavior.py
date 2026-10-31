"""Behaviour model that makes the synthetic data tell a story.

Uniformly random data gives flat dashboards: every cohort retains ~20 %, a
campaign changes nothing, and every game looks alike. This module adds the
patterns real gaming data has, deterministically (same ids -> same traits):

* Players have an activity level, a spending profile (most never pay, a few
  "whales" pay a lot) and a lifetime after which they churn. -> retention decays.
* Games have a popularity (a few hits, a long tail) and a hidden quality that
  drives review ratings. -> top-games charts and rating gaps make sense.
* Sessions follow the Ulaanbaatar day (quiet nights, evening peak).
* Marketing campaigns from data/business/marketing_campaigns.xlsx make the
  promoted game ~2.5x as popular while they run. -> campaign uplift is visible.
* Purchases happen inside sessions, in the player's local currency.
* Reviews and achievements come from player/game pairs that actually played.

Traits are derived from the entity id with a fixed seed, so the initial load and
every later hourly append agree on who is a whale or which game is a hit.
"""

from __future__ import annotations

import csv
import math
import random
from bisect import bisect_right
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Iterable, Sequence

ROOT = Path(__file__).resolve().parents[2]
FX_SEED = ROOT / "dbt" / "seeds" / "fx_rates_to_usd.csv"
CAMPAIGNS_WORKBOOK = ROOT / "data" / "business" / "marketing_campaigns.xlsx"
TRAIT_SEED = "gaming-capstone-traits-v1"

COUNTRY_CURRENCY = {
    "US": "USD", "CA": "CAD", "GB": "GBP", "DE": "EUR", "FR": "EUR",
    "JP": "JPY", "KR": "KRW", "AU": "AUD", "BR": "BRL", "MN": "MNT",
}
# Fallback if the dbt seed is not available (kept identical to it).
DEFAULT_USD_PER_UNIT = {
    "USD": Decimal("1"), "EUR": Decimal("1.08"), "GBP": Decimal("1.27"), "JPY": Decimal("0.0067"),
    "CAD": Decimal("0.73"), "MNT": Decimal("0.00029"), "KRW": Decimal("0.00072"),
    "AUD": Decimal("0.66"), "BRL": Decimal("0.18"),
}
MICROTRANSACTIONS_USD = (Decimal("0.99"), Decimal("1.99"), Decimal("4.99"), Decimal("9.99"))
# Relative session starts by local (UTC+8, no DST) hour: quiet nights, evening peak.
LOCAL_UTC_OFFSET_HOURS = 8
HOURLY_SHAPE = (
    0.55, 0.35, 0.25, 0.20, 0.20, 0.25, 0.40, 0.60, 0.75, 0.85, 0.90, 0.95,
    1.00, 1.05, 1.00, 1.00, 1.10, 1.25, 1.50, 1.75, 1.90, 1.85, 1.55, 1.05,
)
WEEKEND_BOOST = 1.2
CAMPAIGN_BOOST = 2.5  # a game under an active campaign is played ~2.5x as often


@dataclass(frozen=True)
class PlayerTraits:
    activity: float
    spend: float
    lifetime_days: float


@dataclass(frozen=True)
class GameTraits:
    popularity: float
    quality: float


@dataclass(frozen=True)
class Campaign:
    game_id: int
    start: date
    end: date


def _trait_rng(kind: str, entity_id: int) -> random.Random:
    return random.Random(f"{TRAIT_SEED}:{kind}:{entity_id}")


def player_traits(player_id: int) -> PlayerTraits:
    rng = _trait_rng("player", player_id)
    activity = rng.lognormvariate(0.0, 0.8)
    roll = rng.random()
    spend = 0.03 if roll < 0.70 else 1.0 if roll < 0.95 else 8.0
    segment = rng.random()
    if segment < 0.55:            # casual: tries the platform, leaves within weeks
        lifetime = rng.expovariate(1 / 30)
    elif segment < 0.85:          # regular: a few months
        lifetime = rng.expovariate(1 / 150)
    else:                         # loyal: years
        lifetime = rng.uniform(400, 1_500)
    return PlayerTraits(activity=activity, spend=spend, lifetime_days=max(1.0, lifetime))


def game_traits(game_id: int) -> GameTraits:
    rng = _trait_rng("game", game_id)
    popularity = min(rng.paretovariate(1.3), 60.0)
    quality = min(5.0, max(1.0, rng.gauss(3.3, 0.7)))
    return GameTraits(popularity=popularity, quality=quality)


def churn_at(created_at: datetime, traits: PlayerTraits) -> datetime:
    return created_at + timedelta(days=traits.lifetime_days)


def usd_rates(path: Path = FX_SEED) -> dict[str, Decimal]:
    rates = dict(DEFAULT_USD_PER_UNIT)
    if path.exists():
        with path.open(encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                rates[row["currency"].strip()] = Decimal(row["usd_per_unit"].strip())
    return rates


def load_campaigns(path: Path = CAMPAIGNS_WORKBOOK) -> list[Campaign]:
    """Campaign windows from the marketing workbook; [] if it is missing or unreadable."""
    if not path.exists():
        return []
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001 - the generator must work without the workbook
        return []
    if not rows:
        return []
    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    try:
        game_index, start_index, end_index = (
            header.index("game_id"), header.index("start_date"), header.index("end_date")
        )
    except ValueError:
        return []
    campaigns = []
    for row in rows[1:]:
        try:
            start, end = row[start_index], row[end_index]
            start = start.date() if isinstance(start, datetime) else start
            end = end.date() if isinstance(end, datetime) else end
            if isinstance(start, date) and isinstance(end, date) and row[game_index] is not None:
                campaigns.append(Campaign(int(row[game_index]), start, end))
        except (TypeError, ValueError, IndexError):
            continue
    return campaigns


def local_hour_weight(moment: datetime) -> float:
    local = moment + timedelta(hours=LOCAL_UTC_OFFSET_HOURS)
    weight = HOURLY_SHAPE[local.hour] / max(HOURLY_SHAPE)
    return weight * (WEEKEND_BOOST if local.weekday() >= 5 else 1.0) / WEEKEND_BOOST


def activity_time(rng: random.Random, start: datetime, end: datetime, shaped: bool = True) -> datetime:
    """A moment in [start, end]; long spans follow the daily traffic curve (rejection sampling)."""
    span = (end - start).total_seconds()
    if span <= 0:
        return start
    candidate = start
    for _ in range(25):
        candidate = start + timedelta(seconds=rng.uniform(0, span))
        if not shaped or span < 2 * 86_400 or rng.random() < local_hour_weight(candidate):
            return candidate.replace(microsecond=0)
    return candidate.replace(microsecond=0)


class WeightedPicker:
    """Fast weighted choice over a fixed population (cumulative weights + bisect)."""

    def __init__(self, items: Sequence[int], weights: Iterable[float]):
        self.items = list(items)
        self.cumulative: list[float] = []
        total = 0.0
        for weight in weights:
            total += max(weight, 0.0)
            self.cumulative.append(total)
        self.total = total

    def __bool__(self) -> bool:
        return bool(self.items) and self.total > 0

    def pick(self, rng: random.Random) -> int:
        index = bisect_right(self.cumulative, rng.random() * self.total)
        return self.items[min(index, len(self.items) - 1)]


def purchase_amount(rng: random.Random, price_usd: Decimal, currency: str, rates: dict[str, Decimal]) -> Decimal:
    """In-game purchase in local currency: a share of the game's price, or a microtransaction."""
    if price_usd > 0 and rng.random() < 0.6:
        usd = price_usd * Decimal(str(round(rng.uniform(0.5, 1.0), 4)))
    else:
        usd = rng.choice(MICROTRANSACTIONS_USD)
    local = usd / rates.get(currency, Decimal("1"))
    return max(Decimal("0.01"), local.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def review_rating(rng: random.Random, quality: float) -> int:
    return int(min(5, max(1, round(rng.gauss(quality, 0.9)))))


def session_minutes(rng: random.Random) -> int:
    return int(min(360, max(5, math.floor(rng.lognormvariate(3.6, 0.7)))))
